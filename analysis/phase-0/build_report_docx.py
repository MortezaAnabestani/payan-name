import csv
import json
from pathlib import Path

from docx import Document
from docx.enum.section import WD_SECTION_START
from docx.enum.table import WD_ALIGN_VERTICAL, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor

BASE = Path(r"D:\Programming\Me\Payan-name\analysis\phase-0")
OUT = BASE / "report-three-drops.docx"
DATA = json.loads((BASE / "report-data.json").read_text(encoding="utf-8"))
BEATS = list(csv.DictReader((BASE / "narrative-beats.csv").open(encoding="utf-8-sig")))

FONT = "Vazirmatn"
FALLBACK = "Tahoma"
INK = RGBColor(32, 32, 32)
MUTED = RGBColor(90, 90, 90)
ACCENT = RGBColor(30, 55, 85)


def set_run_font(run, size=11, bold=None, color=None, italic=None, ltr=False):
    run.font.name = FONT
    run.font.size = Pt(size)
    if bold is not None:
        run.bold = bold
    if italic is not None:
        run.italic = italic
    if color is not None:
        run.font.color.rgb = color
    rpr = run._element.get_or_add_rPr()
    rfonts = rpr.find(qn("w:rFonts"))
    if rfonts is None:
        rfonts = OxmlElement("w:rFonts")
        rpr.append(rfonts)
    for attr in ("w:ascii", "w:hAnsi", "w:eastAsia", "w:cs"):
        rfonts.set(qn(attr), "Consolas" if ltr else FONT)
    if not ltr and rpr.find(qn("w:rtl")) is None:
        rpr.append(OxmlElement("w:rtl"))


def set_para_rtl(paragraph, align=WD_ALIGN_PARAGRAPH.RIGHT):
    paragraph.alignment = align
    ppr = paragraph._p.get_or_add_pPr()
    if ppr.find(qn("w:bidi")) is None:
        ppr.append(OxmlElement("w:bidi"))


def set_para_ltr(paragraph):
    paragraph.alignment = WD_ALIGN_PARAGRAPH.LEFT
    ppr = paragraph._p.get_or_add_pPr()
    bidi = ppr.find(qn("w:bidi"))
    if bidi is None:
        bidi = OxmlElement("w:bidi")
        ppr.append(bidi)
    bidi.set(qn("w:val"), "0")


def para(text="", size=11, bold=False, color=INK, align=WD_ALIGN_PARAGRAPH.JUSTIFY, space_after=6, style=None, ltr=False):
    p = DOC.add_paragraph(style=style) if style else DOC.add_paragraph()
    if ltr:
        set_para_ltr(p)
    else:
        set_para_rtl(p, align)
    p.paragraph_format.space_after = Pt(space_after)
    p.paragraph_format.line_spacing = 1.18
    if text:
        run = p.add_run(text)
        set_run_font(run, size=size, bold=bold, color=color, ltr=ltr)
    return p


def heading(text, level=1):
    p = DOC.add_paragraph(style=f"Heading {level}")
    set_para_rtl(p, WD_ALIGN_PARAGRAPH.RIGHT)
    p.paragraph_format.keep_with_next = True
    p.paragraph_format.space_before = Pt(16 if level == 1 else 12)
    p.paragraph_format.space_after = Pt(8)
    run = p.add_run(text)
    size = 17 if level == 1 else 13 if level == 2 else 11.5
    set_run_font(run, size=size, bold=True, color=ACCENT)
    return p


def set_table_bidi(table):
    tblpr = table._tbl.tblPr
    if tblpr.find(qn("w:bidiVisual")) is None:
        tblpr.append(OxmlElement("w:bidiVisual"))
    tblw = OxmlElement("w:tblW")
    tblw.set(qn("w:w"), "5000")
    tblw.set(qn("w:type"), "pct")
    tblpr.append(tblw)


def set_cell_margins(table):
    tblpr = table._tbl.tblPr
    mar = OxmlElement("w:tblCellMar")
    for side, width in (("top", 60), ("left", 90), ("bottom", 60), ("right", 90)):
        node = OxmlElement(f"w:{side}")
        node.set(qn("w:w"), str(width))
        node.set(qn("w:type"), "dxa")
        mar.append(node)
    tblpr.append(mar)


def shade(cell, fill):
    tcpr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:fill"), fill)
    tcpr.append(shd)


def no_split(row):
    trpr = row._tr.get_or_add_trPr()
    node = OxmlElement("w:cantSplit")
    trpr.append(node)


def repeat_header(row):
    trpr = row._tr.get_or_add_trPr()
    node = OxmlElement("w:tblHeader")
    trpr.append(node)


def fill_cell(cell, text, size=9.5, bold=False, color=INK, align=WD_ALIGN_PARAGRAPH.RIGHT, ltr=False):
    cell.vertical_alignment = WD_ALIGN_VERTICAL.TOP
    cell.text = ""
    lines = str(text).split("\n") if str(text).strip() else [""]
    for i, line in enumerate(lines):
        p = cell.paragraphs[0] if i == 0 else cell.add_paragraph()
        if ltr:
            set_para_ltr(p)
        else:
            set_para_rtl(p, align)
        p.paragraph_format.space_after = Pt(1)
        p.paragraph_format.space_before = Pt(1)
        p.paragraph_format.line_spacing = 1.05
        run = p.add_run(line)
        set_run_font(run, size=size, bold=bold, color=color, ltr=ltr)


def new_table(cols=3):
    table = DOC.add_table(rows=0, cols=cols)
    table.style = "Table Grid"
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = True
    set_table_bidi(table)
    set_cell_margins(table)
    return table


def section_row(table, text):
    row = table.add_row()
    no_split(row)
    cell = row.cells[0].merge(row.cells[2])
    fill_cell(cell, text, size=10.5, bold=True, color=ACCENT)
    shade(cell, "DCE6F1")


def two_col(table, label, value, label_size=9.5, value_size=9.5, header=False, repeat=False):
    row = table.add_row()
    no_split(row)
    if repeat:
        repeat_header(row)
    fill_cell(row.cells[0], label, size=label_size, bold=True)
    merged = row.cells[1].merge(row.cells[2]) if len(row.cells) == 3 else row.cells[1]
    fill_cell(merged, value, size=value_size)
    if header:
        shade(row.cells[0], "EDEDED")
        shade(merged, "EDEDED")
    else:
        shade(row.cells[0], "F2F2F2")


def three_col(table, a, b, c, header=False, repeat=False):
    row = table.add_row()
    no_split(row)
    if repeat:
        repeat_header(row)
    for cell, text in zip(row.cells, [a, b, c]):
        fill_cell(cell, text, size=9.2, bold=header)
        if header:
            shade(cell, "EDEDED")


def pair_header(table):
    row = table.add_row()
    no_split(row)
    fill_cell(row.cells[0], "چاپ", size=9.2, bold=True)
    merged = row.cells[1].merge(row.cells[2])
    fill_cell(merged, "رایانه", size=9.2, bold=True)
    shade(row.cells[0], "EDEDED")
    shade(merged, "EDEDED")


def pair_row(table, print_text, computer_text):
    row = table.add_row()
    no_split(row)
    fill_cell(row.cells[0], print_text, size=9.2)
    merged = row.cells[1].merge(row.cells[2])
    fill_cell(merged, computer_text, size=9.2)


def level_note(level):
    notes = {
        "۱": "کارکرد به هستهٔ حسی/بدنی رویداد وابسته است؛ پیوند با پروتکل‌های کلان کمتر است.",
        "۲": "کارکرد به قراردادهای کلان روایت، اعتماد، عاملیت یا نهاد وابسته است.",
        "۳": "کارکرد به فرم حوزه‌ای (شعر، موسیقی، زمان/مکان) وابسته است.",
        "۴": "کارکرد به ادراک، فضا، رابط حسی یا کانونی‌سازی وابسته است.",
        "۵": "کارکرد با حفظ هستهٔ انتقال‌پذیر رویداد و بدون واسطهٔ رسانه‌ای سنگین به دست می‌آید.",
        "—": "پیرامتن، بیرون از طیف وابستگی روایی/بدنی واحدهای ۱ تا ۲۷ قرار دارد.",
    }
    return notes.get(level, "سطح اتکاء بر اساس نسبت کارکرد واحد به پروتکل چاپی تعیین شده است.")


def add_unit_table(record):
    form = record["form"]
    heading(f"{record['id']} — {record['title']}", 2)
    meta = (
        f"{record['type']} | {record['chapter']} | سطر {record['source']['lines']} | "
        f"جمله‌های {record['source']['sentences']} | سطح {form['3_level']}"
    )
    para(meta, size=9, color=MUTED, align=WD_ALIGN_PARAGRAPH.RIGHT, space_after=4)

    table = new_table()
    header_row = table.add_row()
    no_split(header_row)
    repeat_header(header_row)
    header_cell = header_row.cells[0].merge(header_row.cells[2])
    fill_cell(
        header_cell,
        f"{record['id']} — {record['title']} | {record['chapter']} | سطح {form['3_level']}",
        size=9,
        bold=True,
        color=ACCENT,
    )
    shade(header_cell, "DCE6F1")
    section_row(table, "۱. واحد پایهٔ چاپی")
    two_col(table, "متن منبع", form["1_print_unit"], value_size=9.2)
    two_col(table, "ارجاع تطبیقی", record["crosswalk"], value_size=9)

    section_row(table, "۲. جایگاه در کلان‌ساختار و روابط بیناواحدی")
    two_col(
        table,
        "موقعیت در توالی روایت",
        f"{form['2_position']['chapter']}؛ ترتیب {form['2_position']['order']}؛ "
        f"سطر {form['2_position']['lines']}؛ جمله‌های {form['2_position']['sentences']}؛ "
        f"بازهٔ کاراکتری {form['2_position']['char_range']}",
    )
    two_col(table, "روابط با واحدهای مجاور", form["2_position"]["relations"])
    two_col(table, "اشتراک موتیفی", "؛ ".join(form["2_position"]["motifs"]))
    two_col(table, "نقش در معماری کلان", form["2_position"]["macro_role"])
    two_col(table, "وارونگی دیالکتیکی بینارسانه‌ای", form["2_position"]["dialectical_inversion"])

    section_row(table, "۳. جانمایی در طیف اتکاء به مسلط چاپی")
    two_col(table, "سطح اتکاء", f"{form['3_level']} — {level_note(form['3_level'])}")

    section_row(table, "۴. مسلط هدف و پروتکل بازتنظیم‌شده")
    three_col(table, "مسلط", "پروتکل طبیعی‌شده", "نحوهٔ بازتنظیم", header=True)
    three_col(
        table,
        form["4_dominant_protocol"]["dominant"],
        form["4_dominant_protocol"]["normalized_protocol"],
        form["4_dominant_protocol"]["recalibration"],
    )

    section_row(table, "۵. کارکرد مسلط‌زدا")
    two_col(table, "کارکرد", form["5_dominant_breaking_function"])

    unit = form["6_transposition_unit"]
    section_row(table, "۶. واحد پایهٔ تراگشت")
    two_col(table, "نام واحد", unit["name"], label_size=9.2, value_size=9.5)
    two_col(table, "سوژه", unit["subject"])
    two_col(table, "پروتکل نقض‌شده", unit["violated_protocol"])
    two_col(table, "مداخله", unit["intervention"])
    two_col(table, "تعلیق", unit["suspension"])
    two_col(table, "نقش بازیکن", unit["player_role"])
    two_col(table, "فعل اصلی", unit["verb"])
    two_col(table, "هدف", unit["goal"])
    two_col(table, "مانع", unit["obstacle"])
    two_col(table, "شرایط شکست", unit["failure"])
    two_col(table, "شرط خروج", unit["exit"])

    section_row(table, "۷. شگرد تراگشت")
    two_col(table, "شگرد نامزد", form["7_technique"]["candidate"])
    two_col(table, "زنجیرهٔ هم‌ارزی", "چاپ: مسلط ← مداخله ← کارکرد | رایانه: پیش‌فرض ← مداخله ← همان کارکرد", value_size=9)
    pair_header(table)
    for pair in form["7_technique"]["print_computer_pairs"]:
        pair_row(table, pair["print"], pair["computer"])
    two_col(table, "کارکرد مشترک", form["12_summary"]["function"], value_size=9)

    section_row(table, "۸. مصالح بومی رایانشی")
    three_col(table, "مصالح بومی", "نحوهٔ ترافکنی کارکرد", "وضعیت", header=True)
    three_col(
        table,
        form["8_native_materials"][0]["material"],
        unit["intervention"],
        "نامزد",
    )
    three_col(
        table,
        form["12_summary"]["materials"],
        form["12_summary"]["function"],
        "نامزد",
    )

    section_row(table, "۹. آستانهٔ تفسیری و حلقهٔ بازخورد زیباشناختی")
    two_col(table, "پیمان تفسیری", form["9_interpretation"]["contract"])
    two_col(table, "آستانهٔ تفسیری", form["9_interpretation"]["threshold"])
    two_col(table, "حلقهٔ بازخورد زیباشناختی", form["9_interpretation"]["feedback_loop"])
    two_col(table, "شرایط شکست تفسیری", "\n".join(form["9_interpretation"]["failure_conditions"]))

    section_row(table, "۱۰. قواعد ترکیب")
    rules = "\n".join(f"{i}. {rule}" for i, rule in enumerate(form["10_composition_rules"], 1))
    two_col(table, "قواعد", rules)

    section_row(table, "۱۱. فضای آزاد")
    two_col(table, "تصمیم‌های تفسیری لازم", "پرسش، گزینه و پیامد هر انتخاب در جدول زیر ثبت شده است.", value_size=9)
    three_col(table, "پرسش", "گزینه", "پیامد", header=True)
    for item in form["11_free_space"]["options"]:
        for index, option in enumerate(item["options"]):
            question = item["question"] if index == 0 else ""
            three_col(table, question, option["option"], option["consequence"])
    two_col(table, "انتقال‌ناپذیری‌های محتمل", form["11_free_space"]["untranslatable"])

    summary = form["12_summary"]
    section_row(table, "۱۲. جمع‌بندی واحد پایه")
    two_col(table, "سطح طیف", summary["level"], label_size=9.2)
    two_col(table, "مسلط هدف", summary["dominant"])
    two_col(table, "پروتکل بازتنظیم‌شده", summary["protocol"])
    two_col(table, "کارکرد مسلط‌زدا", summary["function"])
    two_col(table, "واحد تراگشت", summary["transposition_unit"])
    two_col(table, "شگرد تراگشت", summary["technique"])
    two_col(table, "مصالح بومی", summary["materials"])
    two_col(table, "فضای آزاد", summary["free_space"])

    DOC.add_paragraph()


def add_page_number(section):
    footer = section.footer
    p = footer.paragraphs[0]
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = p.add_run()
    set_run_font(run, size=9, color=MUTED)
    begin = OxmlElement("w:fldChar")
    begin.set(qn("w:fldCharType"), "begin")
    instr = OxmlElement("w:instrText")
    instr.set(qn("xml:space"), "preserve")
    instr.text = " PAGE "
    end = OxmlElement("w:fldChar")
    end.set(qn("w:fldCharType"), "end")
    run._r.append(begin)
    run._r.append(instr)
    run._r.append(end)


DOC = Document()
section = DOC.sections[0]
section.page_width = Cm(21)
section.page_height = Cm(29.7)
section.top_margin = Cm(2.2)
section.bottom_margin = Cm(2)
section.left_margin = Cm(2)
section.right_margin = Cm(2)
section.start_type = WD_SECTION_START.NEW_PAGE
add_page_number(section)

styles = DOC.styles
normal = styles["Normal"]
normal.font.name = FONT
normal.font.size = Pt(11)
normal.element.rPr.rFonts.set(qn("w:cs"), FONT)
normal.paragraph_format.space_after = Pt(6)
normal.paragraph_format.line_spacing = 1.18

for style_name in ("Heading 1", "Heading 2", "Heading 3"):
    style = styles[style_name]
    style.font.name = FONT
    style.element.rPr.rFonts.set(qn("w:cs"), FONT)
    style.font.color.rgb = ACCENT

DOC.core_properties.title = "تراگشت سه قطره خون — گزارش فاز صفر"
DOC.core_properties.subject = "گزارش واحدهای پایه و طراحی تراگشتی"
DOC.core_properties.author = "پایان‌نامه تراگشت ادبی"
DOC.core_properties.keywords = "سه قطره خون، تراگشت، بازی ادبی، رسانه، مسلط‌زدایی"

para("تراگشت سه قطره خون", size=28, bold=True, color=ACCENT, align=WD_ALIGN_PARAGRAPH.CENTER, space_after=8)
para("گزارش واحدهای پایه و طراحی تراگشتی — فاز صفر", size=15, bold=True, color=MUTED, align=WD_ALIGN_PARAGRAPH.CENTER, space_after=28)

meta = new_table(cols=2)
for label, value in [
    ("منبع", "«سه قطره خون» — صادق هدایت"),
    ("واحدهای پایه", "۲۷ واحد + پیرامتن G1"),
    ("بیت‌های L1", "۶۰ بیت"),
    ("وضعیت", "طراحی کامل؛ بازبینی و پلی‌تست انجام نشده"),
    ("تاریخ", "اکتبر ۲۰۲۶"),
]:
    row = meta.add_row()
    no_split(row)
    fill_cell(row.cells[0], label, size=10, bold=True)
    shade(row.cells[0], "F2F2F2")
    fill_cell(row.cells[1], value, size=10)

para("این سند، تحلیل تراگشتی و طراحی ماشین محاسباتی است؛ نه ادعای بازی کامل و نه گزارش پذیرش تجربی.", size=10, color=MUTED, align=WD_ALIGN_PARAGRAPH.CENTER, space_after=4)
para("SHA-256: " + DATA["source_sha256"], size=8.5, color=MUTED, align=WD_ALIGN_PARAGRAPH.CENTER, ltr=True)

DOC.add_page_break()

heading("چکیده", 1)
para("این گزارش، تبدیل کارکرد مسلط‌زدای داستان کوتاه «سه قطره خون» به یک ماشین محاسباتی قابل‌تجربه را در سطح واحدهای پایه طراحی می‌کند. روش کار بر جداسازی دو لایه استوار است: لایهٔ L1، شصت بیت روایی که تمام قصه، جهان و شخصیت را نگه می‌دارد؛ و لایهٔ L2، بیست‌وهفت واحد کارکردی به‌همراه پیرامتن G1 که هر یک یک انتظار طبیعی‌شده را در پروتکل رسانه‌ای نقض می‌کند. هر واحد با نقل‌قول دقیق، بازهٔ کاراکتری، مسلط هدف، پروتکل بازتنظیم‌شده، واحد تراگشت، شگرد، مصالح بومی، پیمان تفسیری، قواعد ترکیب و فضای آزاد مستند شده است. پوشش منبع کامل است: ۱۶۵ جمله، ۱۱٬۱۷۷ نویسهٔ غیرفاصله و ۲۸ نقل‌قول، بدون حذف خاموش. این گزارش وضعیت «طراحی‌شده اما اعتبارنیافته» دارد؛ پلی‌تست، بازبینی انسانی و پذیرش تجربی هنوز انجام نشده است.")
para("واژه‌های کلیدی: تراگشت، مسلط‌زدایی، پروتکل، بازی ادبی، سه قطره خون، واحد تراگشت، فضای آزاد.")

heading("۱. مسئله، پرسش‌ها و روش", 1)
para("پرسش اصلی این نیست که داستان چگونه به چند صحنهٔ قابل‌کلیک ترجمه شود. پرسش این است که کارکرد مسلط‌زدای متن چاپی، یعنی آن لحظه‌ای که انتظار طبیعی‌شدهٔ خواننده از رسانه به هم می‌خورد، چگونه به یک عمل محاسباتی بدل شود که بازیکن آن را انجام دهد، نه فقط بخواند. بنابراین هر واحد باید هم شاهد منبع داشته باشد، هم یک پروتکل طبیعی‌شده را نقض کند، هم در وضعیت بازی پیامد داشته باشد.")
para("روش از تعریف واحد در فصل ۶ رساله آغاز می‌شود: واحد، یک رویداد شناختی‌رویه‌ای با کارکرد مسلط‌زدای مستقل، یا هسته‌ای با کارکردهای وابسته است که باید به یک شیء برنامه‌پذیر بدل شود. این تعریف به دو لایه ترجمه شده است. L1 برای قصه و تداوم؛ L2 برای پروتکل و مداخله. هر L2 از یک یا چند L1 ساخته می‌شود و هر L1 دست‌کم به یک L2 متصل است.")
para("سه پرسش عملیاتی گزارش عبارت‌اند از: آیا هر جملهٔ منبع مالک مشخصی دارد؟ آیا هر واحد یک ماشین حالت، فعل، هدف، مانع و پیامد قابل‌اجرا دارد؟ و آیا هر انتخاب، وضعیت، دانش بازیکن یا رابطه را تغییر می‌دهد، یا فقط یک خط متن را عوض می‌کند؟ جدول‌های بخش ۴ به این پرسش‌ها در سطح هر واحد پاسخ می‌دهند.")
para("این گزارش عمداً از داوری زیبایی‌شناختی نهایی دربارهٔ بازی پرهیز می‌کند. طراحی‌ها «نامزد» هستند و تا زمانی که یک برش عمودی ساخته و پلی‌تست نشود، نمی‌توان دربارهٔ تفسیرپذیری، اصطکاک یا اصالت تجربه ادعای پذیرش کرد.")

heading("۲. تحلیل منبع و پوشش", 1)
para("متن منبع در ۵۷ سطر و چهار بلوک جداشده با ◆ تنظیم شده است. برای پوشش، تمام نویسه‌های غیرفاصله پس از یکسان‌سازی فاصله‌ها و جایگزینی ◆ با فاصله اندازه‌گیری شده‌اند. جدول زیر شاهد عددی پوشش است.")
cov = new_table(cols=2)
two_col(cov, "شاهد", "مقدار", header=True, repeat=True)
for label, value in [
    ("SHA-256", DATA["source_sha256"]),
    ("نویسه‌های غیرفاصله (نرمال‌شده)", f'{DATA["source_chars_normalized"]:,}'),
    ("جمله‌های شناسایی‌شده", "۱۶۵"),
    ("جمله‌های دارای مالک", "۱۶۵ از ۱۶۵"),
    ("نویسه‌های پوشش‌داده‌شده", "۱۱٬۱۷۷ از ۱۱٬۱۷۷"),
    ("نقل‌قول‌های دقیق", "۲۸ از ۲۸"),
    ("هم‌پوشانی کنترل‌شده", "S018 میان U4 و U5؛ مالک: U5"),
    ("حذف خاموش", "صفر"),
]:
    row = cov.add_row()
    no_split(row)
    fill_cell(row.cells[0], label, size=9.5, bold=True)
    shade(row.cells[0], "F2F2F2")
    fill_cell(row.cells[1], value, size=9.5, ltr=(label == "SHA-256"))
para("هم‌پوشانی S018 عمدی است: U4 کاتالیزور خوراک است و درون U5 تودرتو می‌شود؛ مالکیت پوشش با U5 است و U4 فقط بخش نخست جمله را به‌عنوان واحد مستقل ثبت می‌کند. این تصمیم در `coverage-report.md` و `l1-l2-crosswalk.md` مستند شده است.")
para("اطلس L1 از شصت بیت تشکیل شده است. N001 به‌عنوان پیرامتن/عنوان بازتعریف شد؛ N059 که در متن منبع وجود ندارد بایگانی شد؛ N060 وسواس زهر و N061 دعوت/سکوت افزوده شدند. هیچ بیت تازه‌ای به قصه اضافه نشده است؛ افزوده‌ها فقط شکاف‌های پوشش متن منبع را پر می‌کنند.")

heading("۳. نقشهٔ فصل‌ها", 1)
para("چیدمان از توالی منبع پیروی می‌کند. آزادی بازیکن در تفسیر و شهادت است، نه در پریدن از واحدها یا ساختن رخداد تازه. هر فصل با یک وضعیت تفسیری بسته می‌شود، نه با یک پیروزی اخلاقی.")
chap = new_table(cols=3)
three_col(chap, "فصل", "واحدها", "کارکرد", header=True, repeat=True)
for a, b, c in [
    ("پیش‌درآمد", "G1", "نصب موتیف مرکزی به‌مثابۀ دال معلق"),
    ("فصل ۱: کاغذ و آسایشگاه", "U1–U9", "تأسیس جهان بسته، پیوند نوشتن با بدن و بی‌اعتبارکردن تدریجی گزارش راوی"),
    ("فصل ۲: عباس و میهمانان", "U10–U12", "تبدیل هنر به آیین تکرار؛ افشای فرافکنی؛ گذار به انزوای کامل"),
    ("فصل ۳: سیاوش و سه قطره", "U13–U15", "انتقال عاملیت به دعوت سیاوش؛ تبدیل راوی به شاهد؛ قاب اعتراف"),
    ("فصل ۴: نازی", "U16–U20", "جهان ذهنی حیوان؛ میل و خشونت؛ سوگ؛ بازگشت صدا"),
    ("فصل ۵: نمایش عادی‌بودن", "U21–U27", "شهادت کارگردانی‌شده؛ داوری اجتماعی؛ پایان تصویری بدون کاتارسیس"),
]:
    three_col(chap, a, b, c)

heading("۴. جدول‌های واحدهای پایه", 1)
para("هر جدول دوازده بخش فرم نمونه را حفظ می‌کند. بخش‌های ۱ تا ۵ از شاهد منبع و تحلیل مسلط می‌آیند؛ بخش‌های ۶ تا ۱۲ طراحی تراگشت، مصالح، تفسیر، ترکیب، فضای آزاد و جمع‌بندی را ثبت می‌کنند. واحدهای G1 و U1–U27 به ترتیب توالی آمده‌اند.")
para("راهنمای خواندن: «واحد تراگشت» یک ماشین حالت نام‌دار است، نه خلاصهٔ واحد روایی. «فضای آزاد» تصمیم‌هایی است که هنوز به کارگردان/طراح واگذار شده‌اند؛ هر گزینه پیامد متفاوتی دارد و هیچ‌کدام به‌عنوان تصمیم نهایی تثبیت نشده است.", size=10, color=MUTED)

for record in DATA["records"]:
    DOC.add_page_break()
    add_unit_table(record)

heading("۵. معماری فنی و قرارداد پیاده‌سازی", 1)
para("پیاده‌سازی هدف، یک بازی وب با Next.js و TypeScript است. ساختار پیشنهادی، لایه‌های Domain، Application، Ports و Adapters را از رابط جدا می‌کند تا ماشین وضعیت واحدها، دفتر تفسیر و دفتر شهادت، مستقل از React و مرورگر آزمون‌پذیر باشند. رندر متن و UI با DOM/CSS انجام می‌شود؛ Canvas 2D فقط برای واحدهایی که فضای دید/حرکت یا نویز بصری لازم دارند؛ صدا با Web Audio API و لایه‌بندی فضایی. WebGL/3D و حساب کاربری در نسخهٔ اول پیش‌بینی نشده‌اند.")
para("جلسهٔ بازی سرور-محور است: هر فرمان بازیکن به سرور می‌رود، با قواعد Domain اعتبارسنجی می‌شود، رخداد در event log ثبت می‌شود و نمای تازه برمی‌گردد. ذخیره/بارگذاری با رخدادها و snapshotها انجام می‌شود. PostgreSQL + Prisma و Zod نامزدهای فنی‌اند، اما هنوز نصب یا تثبیت نشده‌اند. هیچ حساب کاربری، فروشگاه یا شبکهٔ اجتماعی در فاز اول وجود ندارد.")
para("قرارداد هر واحد شامل شناسه، ترتیب، فصل، نقل‌قول منبع، بازهٔ کاراکتری، نوع، مسلط، پروتکل، سطح، کارکرد، واحد تراگشت، شگرد، مصالح، پیمان تفسیری، قواعد ترکیب، فضای آزاد و پیوندهای گذار است. هر واحد باید یک `UnitDefinition` و یک `UnitView` داشته باشد و فرمان‌های مجازش را اعلام کند. تست‌ها در سه سطح واحد، یکپارچگی جلسه و پیمان تفسیری نوشته می‌شوند.")
para("این بخش، پیش‌نویس معماری است. تصمیم‌های ذخیره‌سازی، مرزهای API و کتابخانه‌های نهایی باید پس از ساخت برش عمودی U1 و اندازه‌گیری تأخیر/هزینه تثبیت شوند.", size=10, color=MUTED)

heading("۶. محدودیت‌ها و وضعیت اعتبار", 1)
para("این گزارش، تحلیل و طراحی است و «بازی کامل» یا «اثر پذیرفته‌شده» نیست. هیچ پلی‌تستی انجام نشده؛ بنابراین ادعای تفسیرپذیری، اصطکاک عامدانه یا اثر عاطفی، اعتبار تجربی ندارد. وضعیت هر واحد `design_complete_unverified` است.")
para("چند انتقال‌ناپذیری مهم ثبت شده‌اند: حس بدنی «بازوی بی‌حس»، طنین فرهنگی «مرغ حق»، بوی گل و هوای باغ، صدای تار دست‌ساز، سوگ حیوانی و ابهام ادبی پایان. در هر مورد، جایگزین رویه‌ای پیشنهاد شده اما افت معنا پذیرفته شده است. این افت‌ها پنهان نشده‌اند.")
para("حق اقتباس و انتشار منبع باید پیش از انتشار عمومی روشن شود. تا آن زمان، این بسته فقط برای تحلیل و نمونه‌سازی داخلی است. محتوای خشونت، حیوان‌آزاری و جنون نیز نیازمند چارچوب هشدار و دسترسی‌پذیری در فاز تولید است.")

heading("پیوست الف: شاهدهای عددی", 1)
app = new_table(cols=2)
two_col(app, "شاهد", "مقدار", header=True, repeat=True)
for label, value in [
    ("منبع", "story.txt"),
    ("SHA-256", DATA["source_sha256"]),
    ("نویسه‌های غیرفاصله", f'{DATA["source_chars_normalized"]:,}'),
    ("جمله‌ها", "۱۶۵"),
    ("واحدهای L2", "۲۷ + G1"),
    ("بیت‌های L1", "۶۰"),
    ("جدول‌های گزارش", "۲۸"),
    ("جفت‌های چاپ/رایانه", "۱۱۲"),
    ("پرسش‌های فضای آزاد", "۸۳"),
]:
    row = app.add_row()
    no_split(row)
    fill_cell(row.cells[0], label, size=9.5, bold=True)
    shade(row.cells[0], "F2F2F2")
    fill_cell(row.cells[1], value, size=9.5, ltr=(label == "SHA-256"))

heading("پیوست ب: فهرست بیت‌های L1", 1)
para("این فهرست برای ردیابی قصه و اتصال L1–L2 آمده است. متن کامل هر بیت در `narrative-beats.csv` و `narrative-beats.md` است.")
beat_table = new_table(cols=3)
three_col(beat_table, "ترتیب", "شناسه / واحد", "خلاصه", header=True, repeat=True)
for beat in BEATS:
    row = beat_table.add_row()
    no_split(row)
    fill_cell(row.cells[0], beat["sequence"], size=8.5)
    fill_cell(row.cells[1], f'{beat["id"]} / {beat["unit_id"]}', size=8.5)
    summary_text = beat.get("summary") or beat.get("beat") or beat.get("description") or ""
    fill_cell(row.cells[2], summary_text, size=8.5)

heading("پیوست ج: فایل‌های داده", 1)
files = new_table(cols=2)
header_row = files.add_row()
no_split(header_row)
repeat_header(header_row)
fill_cell(header_row.cells[0], "فایل", size=9.2, bold=True)
fill_cell(header_row.cells[1], "نقش", size=9.2, bold=True)
shade(header_row.cells[0], "EDEDED")
shade(header_row.cells[1], "EDEDED")
for name, role in [
    ("00-unit-method.md", "روش واحد و قواعد پوشش"),
    ("units-source.json", "منبع واحدهای G1 و U1–U27"),
    ("narrative-units.json", "واحدهای تولیدشده با نقل‌قول و بازه"),
    ("narrative-beats.csv", "اطلس ۶۰ بیت L1"),
    ("l1-l2-crosswalk.md", "اتصال L1–L2"),
    ("01-chapter-map.md", "نقشهٔ فصل‌ها"),
    ("02-architecture.md", "معماری فنی"),
    ("report-data.json", "دادهٔ کامل ۲۸ جدول"),
    ("report-three-drops.docx", "گزارش نهایی Word"),
    ("validate_phase0.py", "اعتبارسنجی بسته"),
]:
    row = files.add_row()
    no_split(row)
    fill_cell(row.cells[0], name, size=8.5, ltr=True)
    fill_cell(row.cells[1], role, size=9)

DOC.save(OUT)
print(f"saved {OUT}")
print(f"tables {len(DOC.tables)} paragraphs {len(DOC.paragraphs)}")
